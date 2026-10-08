import uuid
from typing import List, Optional

import docx
from src.models.document import Document, Formatting, Heading, Paragraph
from src.utils.logging import get_logger

logger = get_logger("ai_writer.parser.docx")

class DocxParser:
    def __init__(self, parserId: str = None):
        self.parserId: str = parserId or str(uuid.uuid4())

    def parse(self, filepath: str) -> Document:
        """Parses a .docx file and returns a populated Document domain model."""
        try:
            doc_obj = docx.Document(filepath)
        except Exception as e:
            logger.error(f"Failed to open {filepath}: {e}")
            raise

        paragraphs: List[Paragraph] = []
        headings: List[Heading] = []
        
        current_pos = 0
        
        for i, docx_p in enumerate(doc_obj.paragraphs):
            text = docx_p.text
            if not text.strip():
                continue
                
            style_name = docx_p.style.name if docx_p.style else "Normal"
            
            formatting = Formatting(
                fontFamily=docx_p.style.font.name if docx_p.style and docx_p.style.font else "Calibri",
                fontSize=docx_p.style.font.size.pt if docx_p.style and docx_p.style.font and docx_p.style.font.size else 11.0,
                alignment=str(docx_p.alignment) if docx_p.alignment else "LEFT"
            )
            
            p_len = len(text)
            
            p = Paragraph(
                id=i,
                text=text,
                startPosition=current_pos,
                endPosition=current_pos + p_len,
                style=style_name,
                formatting=formatting,
                uno_paragraph_ref=docx_p  # Store docx object here for saving later
            )
            paragraphs.append(p)
            
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
                headings.append(Heading(
                    id=i,
                    text=text,
                    level=level,
                    position=current_pos,
                    formatting=formatting
                ))
            
            current_pos += p_len + 1
            
        filename = filepath.split("/")[-1]
        
        domain_doc = Document(
            fileName=filename,
            filePath=filepath,
            fileType="docx",
            paragraphs=paragraphs,
            headings=headings,
            uno_component=doc_obj
        )
        
        domain_doc.create_revision()
        
        return domain_doc
